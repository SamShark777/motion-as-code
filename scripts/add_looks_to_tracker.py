"""One-off: add the "Visual style" column (look per idea), a Looks reference sheet and a
Summary block to shorts/Psychology_Shorts_Tracker.xlsx, keeping its formatting and formulas."""
from copy import copy
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parent.parent
P = ROOT / "shorts" / "Psychology_Shorts_Tracker.xlsx"

LOOKS = [  # id, name, adapted from (lemo-opuscar style), best for, ground hex, accent hex
    ("pixel-lab", "Pixel Lab (house look)", "— (Motion as Code house style)", "Default; social/being-watched ideas", "0A0A0B", "FF4D12"),
    ("whiteboard", "Whiteboard Explainer", "whiteboard", "Learning, study techniques, step-by-step how-tos", "F7F7F4", "2A5CB3"),
    ("data-story", "Data Storytelling", "dataviz", "Ideas whose punch is a number or chart; stats and sampling", "F5F1E8", "B2182B"),
    ("swiss", "Swiss Motion Graphics", "swiss-motion", "Decisions, wording, prices, crisp single-rule ideas", "EAE9E5", "E30613"),
    ("dossier", "Halftone Dossier", "halftone-dossier", "Classic experiments as 'case files', myths on trial", "F1E7D0", "E8417F"),
    ("riso", "Risograph Print", "risograph", "Light, warm, everyday social ideas", "F4EEE2", "FF48B0"),
    ("blueprint", "Blueprint", "blueprint", "Systems and mechanisms: habits, friction, how the brain is 'built'", "18336B", "D6453A"),
    ("terminal", "ASCII / CRT Terminal", "ascii-crt", "Screens, feeds, notifications, attention online", "0B0906", "FFB000"),
    ("keynote", "Dark Tech Keynote", "dark-keynote", "Tech/consumer psychology, product and choice design", "0B0C10", "B7F34A"),
    ("game-show", "Game Show Flat", "game-show", "Quizzes, rewards, competitions, playful demos", "6A4BD8", "FFD23F"),
    ("cartoon-50s", "Mid-century Cartoon", "midcentury-toon", "Practical 'how to' tips, everyday routines", "F2C14E", "E4572E"),
    ("spy-noir", "60s Spy Titles", "spy-titles", "Hidden thoughts, secrets, things people miss or conceal", "1B1714", "E2A52A"),
    ("crayon", "Crayon Picture Book", "crayon-book", "Gentle emotional topics, self-kindness", "FBF5E6", "E2725B"),
    ("pixel-quest", "16-bit Pixel RPG", "pixel-rpg", "Progress, levelling up, motivation, practice", "2B1D4F", "FFD447"),
    ("one-line", "One-line Drawing", "one-line", "Introspective identity, memory and self-image ideas", "F6F2EA", "C8363A"),
]
ASSIGN = {
    1: "pixel-lab", 2: "keynote", 3: "swiss", 4: "dossier", 5: "cartoon-50s", 6: "riso", 7: "spy-noir", 8: "swiss", 9: "keynote", 10: "dossier",
    11: "pixel-lab", 12: "whiteboard", 13: "pixel-quest", 14: "blueprint", 15: "one-line", 16: "dossier", 17: "crayon", 18: "data-story", 19: "riso", 20: "blueprint",
    21: "game-show", 22: "pixel-lab", 23: "spy-noir", 24: "terminal", 25: "riso", 26: "one-line", 27: "swiss", 28: "dossier", 29: "cartoon-50s", 30: "terminal",
    31: "blueprint", 32: "whiteboard", 33: "pixel-quest", 34: "blueprint", 35: "pixel-quest", 36: "crayon", 37: "game-show", 38: "swiss", 39: "cartoon-50s", 40: "data-story",
    41: "crayon", 42: "one-line", 43: "riso", 44: "crayon", 45: "one-line", 46: "whiteboard", 47: "spy-noir", 48: "crayon", 49: "cartoon-50s", 50: "keynote",
    51: "cartoon-50s", 52: "one-line", 53: "spy-noir", 54: "cartoon-50s", 55: "swiss", 56: "crayon", 57: "dossier", 58: "game-show", 59: "riso", 60: "terminal",
    61: "dossier", 62: "data-story", 63: "game-show", 64: "riso", 65: "dossier", 66: "blueprint", 67: "game-show", 68: "spy-noir", 69: "pixel-quest", 70: "data-story",
    71: "data-story", 72: "game-show", 73: "spy-noir", 74: "whiteboard", 75: "one-line", 76: "dossier", 77: "swiss", 78: "one-line", 79: "pixel-quest", 80: "spy-noir",
    81: "terminal", 82: "game-show", 83: "keynote", 84: "riso", 85: "terminal", 86: "swiss", 87: "dossier", 88: "keynote", 89: "cartoon-50s", 90: "crayon",
    91: "whiteboard", 92: "blueprint", 93: "dossier", 94: "spy-noir", 95: "game-show", 96: "one-line", 97: "data-story", 98: "data-story", 99: "dossier", 100: "whiteboard",
}

wb = openpyxl.load_workbook(P)
ws = wb["Shorts ideas"]
HR, R0, R1 = 5, 6, 105
H = {ws.cell(HR, c).value: c for c in range(1, ws.max_column + 1)}
if "Visual style" in H:
    col = H["Visual style"]
else:
    col = ws.max_column + 1
    for src_row, dst_row in ((HR, HR),):
        s, d = ws.cell(src_row, col - 1), ws.cell(dst_row, col)
        d.font, d.fill, d.alignment = copy(s.font), copy(s.fill), copy(s.alignment)
    ws.cell(HR, col).value = "Visual style"
    g = ws.cell(4, col); g.value = "STYLE"; g.font = Font(name="Arial", size=10, bold=True, color="FF555555")
    from openpyxl.utils import get_column_letter
    ws.column_dimensions[get_column_letter(col)].width = 15
    ws.cell(HR, col).comment = Comment("The look this Short is built in (see the Looks sheet). Change it before the idea is made to "
                                       "switch styles; the daily pipeline copies it into the script.", "Pipeline")
for r in range(R0, R1 + 1):
    no = ws.cell(r, 1).value
    cell = ws.cell(r, col)
    src = ws.cell(r, col - 1)
    cell.font = Font(name="Arial", size=11, color="FF0000FF")
    cell.alignment = copy(src.alignment)
    status = ws.cell(r, H["Status"]).value
    # rows already made keep the look they were actually built in
    cell.value = "pixel-lab" if status in ("Scheduled", "Published") else ASSIGN[no]
from openpyxl.utils import get_column_letter as L
dv = DataValidation(type="list", formula1=f"=Looks!$A$5:$A${4 + len(LOOKS)}", allow_blank=True)
dv.add(f"{L(col)}{R0}:{L(col)}{R1}")
ws.add_data_validation(dv)

# Looks sheet
if "Looks" in wb.sheetnames:
    del wb["Looks"]
S = wb.create_sheet("Looks", 1)
S["A1"].value = "Visual styles"
S["A1"].font = Font(name="Arial", size=16, bold=True, color="FF111111")
S["A2"].value = ("Each Short is built in one look. Looks adapt palettes, type and conventions of the lemo-opuscar film styles "
                 "(LemoLab, MIT licence, github.com/lemomo-ai/lemo-opuscar) to the vertical psychology template.")
S["A2"].font = Font(name="Arial", size=10, color="FF555555")
hdr = ws.cell(HR, 1)
for j, h in enumerate(["Look id", "Name", "Adapted from", "Best for", "Ideas assigned", "Ground", "Accent"]):
    c = S.cell(4, 1 + j, h)
    c.font, c.fill, c.alignment = copy(hdr.font), copy(hdr.fill), copy(hdr.alignment)
for i, (lid, name, frm, best, g, a) in enumerate(LOOKS):
    r = 5 + i
    for j, v in enumerate([lid, name, frm, best]):
        S.cell(r, 1 + j, v).font = Font(name="Arial", size=11)
    S.cell(r, 5, f"=COUNTIF('Shorts ideas'!${L(col)}${R0}:${L(col)}${R1},A{r})").font = Font(name="Arial", size=11)
    S.cell(r, 6).fill = PatternFill("solid", fgColor="FF" + g)
    S.cell(r, 7).fill = PatternFill("solid", fgColor="FF" + a)
for c_, w in zip("ABCDEFG", (14, 26, 30, 58, 14, 9, 9)):
    S.column_dimensions[c_].width = w

# Summary block by visual style
Sm = wb["Summary"]
if not any(Sm.cell(r, 1).value == "Visual style" for r in range(1, Sm.max_row + 1)):
    top = Sm.max_row + 2
    heads = ["Visual style", "Published", "Avg views d7", "Avg retention d7", "Avg engagement d7", "Avg views d30"]
    ref = Sm.cell(4, 1)
    for j, h in enumerate(heads):
        c = Sm.cell(top, 1 + j, h)
        c.font, c.fill, c.alignment = copy(ref.font), copy(ref.fill), copy(ref.alignment)
    SH = "'Shorts ideas'"
    hcol = {n: L(H[n]) for n in ("Status", "Total views d7", "YT retention d7", "Engagement d7", "Total views d30")}
    rng = lambda n: f"{SH}!${hcol[n]}${R0}:${hcol[n]}${R1}"
    key = f"{SH}!${L(col)}${R0}:${L(col)}${R1}"
    for i, (lid, *_rest) in enumerate(LOOKS, start=1):
        rr = top + i
        Sm.cell(rr, 1, lid)
        Sm.cell(rr, 2, f'=COUNTIFS({key},$A{rr},{rng("Status")},"Published")')
        Sm.cell(rr, 3, f'=IFERROR(AVERAGEIFS({rng("Total views d7")},{key},$A{rr},{rng("Total views d7")},">0"),"")').number_format = "#,##0"
        Sm.cell(rr, 4, f'=IFERROR(AVERAGEIFS({rng("YT retention d7")},{key},$A{rr},{rng("YT retention d7")},">0"),"")').number_format = "0.0%"
        Sm.cell(rr, 5, f'=IFERROR(AVERAGEIFS({rng("Engagement d7")},{key},$A{rr},{rng("Engagement d7")},">0"),"")').number_format = "0.0%"
        Sm.cell(rr, 6, f'=IFERROR(AVERAGEIFS({rng("Total views d30")},{key},$A{rr},{rng("Total views d30")},">0"),"")').number_format = "#,##0"
        for j in range(1, 7):
            Sm.cell(rr, j).font = Font(name="Arial", size=11)
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(P)
print("Visual style column", L(col))
