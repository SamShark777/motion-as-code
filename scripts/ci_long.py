"""Long-form Ted type 1 (16:9, ~8–10 min, YouTube) — CI steps for .github/workflows/longform.yml on main.
Queue: the "Long form" sheet of shorts/Psychology_Shorts_Tracker.xlsx; scripts in shorts/long/<slug>/script.json.

    python3 scripts/ci_long.py plan [slug]        # pick the next Scripted row (or the given slug) -> $GITHUB_OUTPUT slug/date
    python3 scripts/ci_long.py chunks             # after prep: split the voiceover into render chunks -> $GITHUB_OUTPUT matrix
    python3 scripts/ci_long.py finalize <slug> <date>   # concat chunks, music, mux, size-fit, thumbnail, chapters, host, Buffer, tracker, issue.md
"""
import datetime as dt, json, os, re, shutil, subprocess, sys, tempfile, time
from pathlib import Path
from zoneinfo import ZoneInfo
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
TZ = ZoneInfo("America/Los_Angeles")
REPO = os.environ.get("GITHUB_REPOSITORY", "SamShark777/motion-as-code")
XLSX = ROOT / "shorts" / "Psychology_Shorts_Tracker.xlsx"
POST_AT = os.environ.get("LONG_POST_AT", "12:00")
CHUNK_S = 75
MAX_MB = 92


def out(k, v):
    p = os.environ.get("GITHUB_OUTPUT")
    if p:
        with open(p, "a") as f:
            f.write(f"{k}={v}\n")
    print(f"{k}={v}")


def sheet():
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Long form"]
    H = {ws.cell(4, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(4, c).value}
    rows = {ws.cell(r, 1).value: r for r in range(5, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, int)}
    return wb, ws, H, rows


def set_row(no, **kv):
    from openpyxl.workbook.properties import CalcProperties
    wb, ws, H, rows = sheet()
    r = rows[no]
    for k, v in kv.items():
        ws.cell(r, H[k]).value = v
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(XLSX)


def plan(slug=None):
    wb, ws, H, rows = sheet()
    day = dt.datetime.now(TZ).date() + dt.timedelta(days=1)
    pick = None
    for no, r in sorted(rows.items(), key=lambda kv: (ws.cell(kv[1], H["Priority"]).value or 999, kv[0])):
        s = ws.cell(r, H["Slug"]).value
        if slug and s == slug or (not slug and ws.cell(r, H["Status"]).value == "Scripted" and s
                                  and (ROOT / "shorts" / "long" / s / "script.json").exists()):
            pick = (no, s)
            break
    if not pick and slug and (ROOT / "shorts" / "long" / slug / "script.json").exists():
        pick = (0, slug)
    if not pick:
        out("slug", "")
        return
    out("slug", pick[1]); out("no", pick[0]); out("date", day.isoformat())


def chunks():
    import soundfile as sf
    info = sf.info(str(ROOT / "audio" / "voiceover.wav"))
    d = info.frames / info.samplerate
    n = max(1, int(-(-d // CHUNK_S)))
    m = [{"k": k, "from": k * CHUNK_S, "to": min(d, (k + 1) * CHUNK_S)} for k in range(n)]
    out("matrix", json.dumps({"include": m}))
    out("duration", round(d, 2))


def sh(*a, **k):
    return subprocess.run(a, check=True, text=True, capture_output=True, **k).stdout


def fmt_ts(s):
    s = int(s)
    return f"{s // 60}:{s % 60:02d}" if s < 3600 else f"{s // 3600}:{s % 3600 // 60:02d}:{s % 60:02d}"


def chapters(S):
    ly = json.loads((ROOT / "data" / "lyrics.json").read_text())["lines"]
    rows = []
    for i, ln in enumerate(S["lines"]):
        if ln.get("chapter") and i < len(ly):
            rows.append((0 if not rows else ly[i]["start"], ln["chapter"]))
    if len(rows) < 3:
        return ""
    return "\n".join(f"{fmt_ts(t)} {c}" for t, c in rows)


def thumbnail(S, video, dst):
    """1280x720: a frame from the video on the left, big 2–4 word title on the right."""
    from PIL import Image, ImageDraw, ImageFont, ImageFilter
    t = float(S.get("thumb_at") or 20)
    png = subprocess.run(["ffmpeg", "-v", "error", "-ss", str(t), "-i", str(video), "-frames:v", "1", "-f", "image2pipe", "-vcodec", "png", "-"], capture_output=True).stdout
    import io
    fr = Image.open(io.BytesIO(png)).convert("RGB").resize((1280, 720))
    im = Image.new("RGB", (1280, 720), "#0b0b0d")
    card = fr.crop((240, 0, 1040, 720)).resize((720, 648))
    im.paste(card, (40, 36))
    d = ImageDraw.Draw(im)
    words = (S.get("thumb_text") or " ".join(S["title"])).upper().split()
    fpath = ROOT / "app" / "public" / "fonts" / "ext"
    font_file = next((str(p) for p in fpath.glob("Inter-800*")), None) or next((str(p) for p in fpath.glob("*800*")), None)
    lines, cur = [], ""
    for w in words:
        if len(cur + " " + w) > 9 and cur:
            lines.append(cur); cur = w
        else:
            cur = (cur + " " + w).strip()
    lines.append(cur)
    size = 92 if len(lines) <= 3 else 72
    F = None
    try:
        if font_file and font_file.endswith(".woff2"):
            from fontTools.ttLib import TTFont
            f = TTFont(font_file); f.flavor = None
            ttf = Path(tempfile.mkdtemp()) / "thumb.ttf"; f.save(str(ttf)); font_file = str(ttf)
        F = ImageFont.truetype(font_file, size)
    except Exception as e:
        print("thumbnail font fallback:", e)
        for cand in ("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",):
            if Path(cand).exists():
                F = ImageFont.truetype(cand, size)
        F = F or ImageFont.load_default(size)
    y = 360 - len(lines) * size * 0.58
    for i, ln in enumerate(lines):
        col = "#FFD23F" if i == len(lines) - 1 else "#FFFFFF"
        d.text((800, y), ln, font=F, fill=col, stroke_width=6, stroke_fill="#000000")
        y += size * 1.12
    im.save(dst, quality=90)


def host(files, date, slug):
    tmp = Path(tempfile.mkdtemp())
    url = f"https://x-access-token:{os.environ['GITHUB_TOKEN']}@github.com/{REPO}.git"
    keep_from = (dt.datetime.now(TZ).date() - dt.timedelta(days=2)).isoformat()
    try:
        sh("git", "clone", "-q", "--depth", "1", "--branch", "media", url, str(tmp / "old"))
        for p in (tmp / "old").iterdir():
            if re.fullmatch(r"\d{4}-\d\d-\d\d", p.name) and p.name >= keep_from:
                shutil.copytree(p, tmp / "new" / p.name)
    except subprocess.CalledProcessError:
        pass
    dd = tmp / "new" / date
    dd.mkdir(parents=True, exist_ok=True)
    for f in files:
        shutil.copy(f, dd / Path(f).name)
    (tmp / "new" / "README.md").write_text("Hosting for scheduled videos (Buffer fetches them from here). Rewritten on every run.\n")
    g = ["git", "-C", str(tmp / "new")]
    sh(*g, "init", "-q"); sh(*g, "checkout", "-q", "-b", "media"); sh(*g, "add", ".")
    sh(*g, "-c", "user.name=shorts-bot", "-c", "user.email=shorts-bot@users.noreply.github.com", "commit", "-qm", f"media: long {slug}")
    sh(*g, "push", "-q", "-f", url, "media")
    return [f"https://raw.githubusercontent.com/{REPO}/media/{date}/{Path(f).name}" for f in files]


def finalize(slug, date, no):
    S = json.loads((ROOT / "shorts" / "long" / slug / "script.json").read_text())
    parts = sorted(Path("chunks").glob("chunk_*.mp4"), key=lambda p: int(re.search(r"(\d+)", p.stem).group(1)))
    lst = Path("chunks/list.txt")
    lst.write_text("".join(f"file '{p.name}'\n" for p in parts))
    sh("ffmpeg", "-v", "error", "-y", "-f", "concat", "-safe", "0", "-i", str(lst), "-c", "copy", "out/long.raw.mp4")
    sh(sys.executable, str(ROOT / "analysis" / "ted_music.py"), "--level", "-24")
    final = Path("out") / f"{slug}.mp4"
    dur = float(sh("ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", "out/long.raw.mp4"))
    sh("ffmpeg", "-v", "error", "-y", "-i", "out/long.raw.mp4", "-i", "out/mix.wav", "-map", "0:v", "-map", "1:a", "-c:v", "libx264", "-preset", "medium",
       "-crf", "24", "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "160k", "-shortest", "-movflags", "+faststart", str(final))
    mb = final.stat().st_size / 1e6
    if mb > MAX_MB:  # GitHub's 100 MB file limit on the media branch
        kbps = int((MAX_MB * 8000 / dur) * 0.97) - 160
        sh("ffmpeg", "-v", "error", "-y", "-i", "out/long.raw.mp4", "-i", "out/mix.wav", "-map", "0:v", "-map", "1:a", "-c:v", "libx264", "-preset", "medium",
           "-b:v", f"{kbps}k", "-maxrate", f"{int(kbps * 1.4)}k", "-bufsize", f"{kbps * 2}k", "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "160k", "-shortest",
           "-movflags", "+faststart", str(final))
        mb = final.stat().st_size / 1e6
    thumb = Path("out") / f"{slug}.jpg"
    thumbnail(S, final, thumb)
    chap = chapters(S)
    vurl, turl = host([final, thumb], date, slug)
    for _ in range(24):
        if subprocess.run(["curl", "-sIL", "-o", "/dev/null", "-w", "%{http_code}", vurl], capture_output=True, text=True).stdout == "200":
            break
        time.sleep(5)
    post = S["post"]
    desc = post.get("description", "")
    if chap:
        desc = f"{desc}\n\nChapters\n{chap}"
    due = dt.datetime.combine(dt.date.fromisoformat(date), dt.time.fromisoformat(POST_AT), TZ).astimezone(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")
    dry = not os.environ.get("BUFFER_API_KEY") or os.environ.get("DRY_RUN") == "1"
    lines = [f"**Long-form for {date} {POST_AT} PT** — {post['title']}", "", f"{dur / 60:.1f} min · {mb:.0f} MB", "",
             f"![thumbnail]({turl})", "", f"[video]({vurl})", ""]
    if chap:
        lines += ["Chapters:", "```", chap, "```", ""]
    if dry:
        lines.insert(1, "🧪 **Dry run — not scheduled**")
    else:
        r = json.loads(sh(sys.executable, str(ROOT / "scripts" / "buffer.py"), "post-long", slug, due, vurl, turl, desc))
        if r.get("youtube"):
            lines.insert(1, "✅ **Scheduled on YouTube** — to veto, delete it in [Buffer](https://publish.buffer.com) before then.")
            if no:
                set_row(int(no), **{"Status": "Scheduled", "Post date": dt.date.fromisoformat(date), "Title used": post["title"],
                                    "Length (min)": round(dur / 60, 1), "Lesson / notes": f"buffer:youtube={r['youtube']}"})
        else:
            lines.insert(1, f"❌ **Buffer rejected it** — {json.dumps(r.get('errors'))}")
            if no:
                set_row(int(no), **{"Lesson / notes": f"Buffer rejected {dt.date.today()}: {json.dumps(r.get('errors'))[:200]}"})
    Path("issue.md").write_text("\n".join(lines))
    print("\n".join(lines))


if __name__ == "__main__":
    c = sys.argv[1] if len(sys.argv) > 1 else ""
    if c == "plan":
        plan(sys.argv[2] if len(sys.argv) > 2 and sys.argv[2] else None)
    elif c == "chunks":
        chunks()
    elif c == "finalize":
        finalize(sys.argv[2], sys.argv[3], sys.argv[4] if len(sys.argv) > 4 else 0)
    else:
        sys.exit(__doc__)
