"""One-off (safe to re-run): add the STORY columns to the "Shorts ideas" sheet of
shorts/Psychology_Shorts_Tracker.xlsx, keeping its formatting and formulas. A row with Format=story is
made as a story Short (see shorts/RUNBOOK.md, "Story Shorts"); the other columns describe the video.
    python3 scripts/add_story_columns.py"""
from copy import copy
from pathlib import Path
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

P = Path(__file__).resolve().parent.parent / "shorts" / "Psychology_Shorts_Tracker.xlsx"
HR, R0 = 5, 6
COLS = [  # header, width, help (cell comment)
    ("Format", 10, "ted2 (default when empty) or story. 'story' makes a fully art-directed Short from the columns "
                   "to the right: Claude writes the narration, the scene-by-scene animation and the sound cues from them."),
    ("Animation style", 34, "How it should look and move, e.g. 'flat neon editorial like One More Never Enough: one "
                            "curly-haired character, bold outlines, camera push-ins, card-swipe transitions'."),
    ("Storyboard", 60, "Optional scene-by-scene plan, one scene per line: what we see, what is said, the headline. "
                       "e.g. '1. Hook: thumb flicks cards on a phone; pull back to reveal the face — VO: You opened "
                       "your phone for a minute.' Empty: Claude plans the scenes from the hook and outline."),
    ("Palette & type", 30, "Optional colours (hex or words) and fonts (any Google Font), e.g. 'white page, ink "
                           "#210B3D, violet #7624CE, neon green highlight; Barlow Condensed headlines'."),
    ("Characters", 26, "Optional: who appears, e.g. 'the One More protagonist (curly hair, violet jacket)', or 'none — objects only'."),
    ("Sound cues", 34, "Optional: sound effects and moments, e.g. 'soft swipe on each card, notification ping on the cue, "
                       "three wooden taps on the end icons'. Available: swipe, whoosh, pop, tick, chime, chime-bright, "
                       "rising, notification, flip, flip-dull, lightning, airy, card-slide, settle, toggle, drawer, "
                       "page-turn, wood-tap. Music: the soft F-major pad bed."),
]

wb = openpyxl.load_workbook(P)
ws = wb["Shorts ideas"]
H = {ws.cell(HR, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(HR, c).value}
last_row = max(r for r in range(R0, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, int))
ref = ws.max_column  # copy header/body styling from the current last column
added = []
for name, width, help_ in COLS:
    if name in H:
        continue
    c = ws.max_column + 1
    s, d = ws.cell(HR, ref), ws.cell(HR, c)
    d.font, d.fill, d.alignment, d.border = copy(s.font), copy(s.fill), copy(s.alignment), copy(s.border)
    d.value = name
    d.comment = Comment(help_, "Pipeline", width=420, height=160)
    ws.column_dimensions[get_column_letter(c)].width = width
    for r in range(R0, last_row + 1):
        src, cell = ws.cell(r, ref), ws.cell(r, c)
        cell.font = Font(name="Arial", size=11, color="FF0000FF")
        cell.border, cell.fill = copy(src.border), copy(src.fill)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    H[name] = c
    added.append(name)
if added:
    g = ws.cell(4, H[COLS[0][0]])
    g.value = "STORY"
    g.font = Font(name="Arial", size=10, bold=True, color="FF555555")
    fc = H["Format"]
    dv = DataValidation(type="list", formula1='"ted2,story"', allow_blank=True)
    dv.add(f"{get_column_letter(fc)}{R0}:{get_column_letter(fc)}{last_row}")
    ws.add_data_validation(dv)
    wb.save(P)
print("added:", ", ".join(added) or "nothing (already there)")
