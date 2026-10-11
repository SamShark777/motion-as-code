"""One-off: turn Shivam's 100-idea sheet into the pipeline tracker.
Keeps the original sheet's columns and style; adds production + performance columns (d7 / d14 / d30),
formulas for totals, retention, engagement and a verdict vs the channel median, and a Summary sheet.
    python3 scripts/make_tracker.py <input.xlsx> <output.xlsx>
"""
import sys
from copy import copy
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)
ws = wb["Shorts ideas"]
HR, R0, R1 = 5, 6, 105  # header row, first/last idea row

hdr_font, hdr_fill, hdr_al = copy(ws["H5"].font), copy(ws["H5"].fill), copy(ws["H5"].alignment)
body_font, body_al = copy(ws["H6"].font), copy(ws["H6"].alignment)
grp_font = Font(name="Arial", size=10, bold=True, color="FF555555")
calc_font = Font(name="Arial", size=11, color="FF111111")          # formulas: black
input_font = Font(name="Arial", size=11, color="FF0000FF")         # filled by the pipeline: blue

# (header, width, kind, group)  kind: input | formula
COLS = [
    ("Priority", 9, "input", "PRODUCTION"), ("Template fit", 12, "input", None), ("Slug", 24, "input", None),
    ("Post date", 12, "input", None), ("Weekday", 10, "formula", None), ("Hook style", 11, "input", None),
    ("Length (s)", 10, "input", None), ("Title used", 40, "input", None), ("Source checked", 30, "input", None),
    ("Metricool ID", 14, "input", None),
    ("YT views d7", 11, "input", "DAY 7"), ("YT avg view (s) d7", 12, "input", None), ("YT retention d7", 12, "formula", None),
    ("YT likes d7", 10, "input", None), ("YT comments d7", 11, "input", None), ("YT shares d7", 10, "input", None),
    ("IG views d7", 11, "input", None), ("IG likes d7", 10, "input", None), ("TT views d7", 11, "input", None),
    ("TT likes d7", 10, "input", None), ("Total views d7", 12, "formula", None), ("Engagement d7", 12, "formula", None),
    ("YT views d14", 11, "input", "DAY 14"), ("IG views d14", 11, "input", None), ("TT views d14", 11, "input", None),
    ("Total views d14", 12, "formula", None),
    ("YT views d30", 11, "input", "DAY 30"), ("IG views d30", 11, "input", None), ("TT views d30", 11, "input", None),
    ("Total views d30", 12, "formula", None),
    ("Verdict d7", 11, "formula", "LEARNING"), ("Lesson / notes", 50, "input", None),
]
start = 9  # column I
col = {name: L(start + i) for i, (name, *_) in enumerate(COLS)}
for i, (name, w, kind, grp) in enumerate(COLS):
    c = L(start + i)
    ws.column_dimensions[c].width = w
    h = ws[f"{c}{HR}"]
    h.value, h.font, h.fill, h.alignment = name, copy(hdr_font), copy(hdr_fill), copy(hdr_al)
    if grp:
        g = ws[f"{c}4"]
        g.value, g.font = grp, grp_font

# per-row formulas
for r in range(R0, R1 + 1):
    c = lambda n: f"{col[n]}{r}"
    f = {
        "Weekday": f'=IF({c("Post date")}="","",TEXT({c("Post date")},"ddd"))',
        "YT retention d7": f'=IF(OR({c("YT avg view (s) d7")}="",N({c("Length (s)")})=0),"",{c("YT avg view (s) d7")}/{c("Length (s)")})',
        "Total views d7": f'=IF(COUNT({c("YT views d7")},{c("IG views d7")},{c("TT views d7")})=0,"",SUM({c("YT views d7")},{c("IG views d7")},{c("TT views d7")}))',
        "Engagement d7": f'=IF(N({c("Total views d7")})=0,"",SUM({c("YT likes d7")},{c("YT comments d7")},{c("YT shares d7")},{c("IG likes d7")},{c("TT likes d7")})/{c("Total views d7")})',
        "Total views d14": f'=IF(COUNT({c("YT views d14")},{c("IG views d14")},{c("TT views d14")})=0,"",SUM({c("YT views d14")},{c("IG views d14")},{c("TT views d14")}))',
        "Total views d30": f'=IF(COUNT({c("YT views d30")},{c("IG views d30")},{c("TT views d30")})=0,"",SUM({c("YT views d30")},{c("IG views d30")},{c("TT views d30")}))',
        "Verdict d7": (f'=IF(N({c("Total views d7")})=0,"",IF(COUNT(${col["Total views d7"]}${R0}:${col["Total views d7"]}${R1})<5,"Too early",'
                       f'IF({c("Total views d7")}>=1.5*MEDIAN(${col["Total views d7"]}${R0}:${col["Total views d7"]}${R1}),"Top",'
                       f'IF({c("Total views d7")}<=0.5*MEDIAN(${col["Total views d7"]}${R0}:${col["Total views d7"]}${R1}),"Low","Average"))))'),
    }
    for i, (name, _, kind, _) in enumerate(COLS):
        cell = ws[f"{L(start + i)}{r}"]
        cell.alignment = copy(body_al)
        cell.font = copy(calc_font if kind == "formula" else input_font)
        if name in f:
            cell.value = f[name]
    ws[f"{col['Post date']}{r}"].number_format = "yyyy-mm-dd"
    for n in ("YT retention d7", "Engagement d7"):
        ws[f"{col[n]}{r}"].number_format = "0.0%"
    for n in ("YT views d7", "IG views d7", "TT views d7", "Total views d7", "YT views d14", "IG views d14", "TT views d14",
              "Total views d14", "YT views d30", "IG views d30", "TT views d30", "Total views d30"):
        ws[f"{col[n]}{r}"].number_format = "#,##0"

# ---- initial plan: template fit + priority (round-robin across themes, strong fits first)
WEAK = {21, 23, 26, 27, 31, 33, 34, 38, 39, 40, 41, 44, 45, 46, 48, 49, 50, 51, 52, 53, 54, 55, 56, 57, 60, 78, 79, 81, 100}
rows = {ws[f"A{r}"].value: r for r in range(R0, R1 + 1)}
themes = []
for no in range(1, 101):
    t = ws[f"B{rows[no]}"].value
    if t not in themes:
        themes.append(t)
for no, r in rows.items():
    ws[f"{col['Template fit']}{r}"].value = "Weak" if no in WEAK else "Strong"
order = []
for tier in (False, True):
    queues = [[no for no in range(1, 101) if ws[f"B{rows[no]}"].value == t and (no in WEAK) == tier] for t in themes]
    while any(queues):
        for q in queues:
            if q:
                order.append(q.pop(0))
for p, no in enumerate(order, start=1):
    ws[f"{col['Priority']}{rows[no]}"].value = p

# idea #1 already went out (horizontal test, YouTube only)
r = rows[1]
ws[f"H{r}"].value = "Published"
ws[f"{col['Slug']}{r}"].value = "spotlight-effect"
import datetime
ws[f"{col['Post date']}{r}"].value = datetime.date(2026, 10, 11)
ws[f"{col['Hook style']}{r}"].value = "eyes"
ws[f"{col['Length (s)']}{r}"].value = 27
ws[f"{col['Title used']}{r}"].value = "Why It Feels Like Everyone's Watching You | The Spotlight Effect"
ws[f"{col['Source checked']}{r}"].value = "Gilovich, Medvec & Savitsky (2000), JPSP 78(2)"
ws[f"{col['Metricool ID']}{r}"].value = "393438157"
ws[f"{col['Lesson / notes']}{r}"].value = "Pilot: 16:9 horizontal, YouTube only — not comparable with the vertical Shorts."

# status list: keep the original values, add the pipeline's
ws.data_validations.dataValidation = []
dv = DataValidation(type="list", formula1='"Idea,Researching,Scripted,Scheduled,Published,Vetoed,Failed,Filming,Editing"', allow_blank=True)
dv.add(f"H{R0}:H{R1}")
ws.add_data_validation(dv)
dv2 = DataValidation(type="list", formula1='"Strong,Weak"', allow_blank=True)
dv2.add(f"{col['Template fit']}{R0}:{col['Template fit']}{R1}")
ws.add_data_validation(dv2)

# verdict colours
vr = f"{col['Verdict d7']}{R0}:{col['Verdict d7']}{R1}"
ws.conditional_formatting.add(vr, CellIsRule(operator="equal", formula=['"Top"'], fill=PatternFill("solid", fgColor="FFD9EAD3")))
ws.conditional_formatting.add(vr, CellIsRule(operator="equal", formula=['"Low"'], fill=PatternFill("solid", fgColor="FFF4CCCC")))

ws.freeze_panes = "D6"
ws[f"{col['Priority']}4"].comment = Comment("Blue = filled in by the daily pipeline (you can edit too). Black = formulas. "
                                            "Priority: lowest number with Status 'Idea' is made next — renumber to reorder.", "Pipeline")
ws["A3"].value = (ws["A3"].value or "") + "  Columns I→AN are filled by the daily Shorts pipeline; Priority sets the order."

# ---- Summary sheet
S = wb.create_sheet("Summary")
S["A1"].value = "What's working"
S["A1"].font = Font(name="Arial", size=16, bold=True, color="FF111111")
S["A2"].value = "Averages over published Shorts with day-7 numbers. Updates as the pipeline fills the tracker."
S["A2"].font = Font(name="Arial", size=10, color="FF555555")
SH = "'Shorts ideas'"
rng = lambda n: f"{SH}!${col[n]}${R0}:${col[n]}${R1}"
def block(top, title, key_range, keys):
    heads = [title, "Published", "Avg views d7", "Avg retention d7", "Avg engagement d7", "Avg views d30"]
    for j, h in enumerate(heads):
        c = S.cell(row=top, column=1 + j, value=h)
        c.font, c.fill, c.alignment = copy(hdr_font), copy(hdr_fill), copy(hdr_al)
    for i, k in enumerate(keys, start=1):
        rr = top + i
        S.cell(row=rr, column=1, value=k).font = Font(name="Arial", size=11)
        a = f"$A{rr}"
        S.cell(row=rr, column=2, value=f'=COUNTIFS({key_range},{a},{SH}!$H${R0}:$H${R1},"Published")')
        S.cell(row=rr, column=3, value=f'=IFERROR(AVERAGEIFS({rng("Total views d7")},{key_range},{a},{rng("Total views d7")},">0"),"")').number_format = "#,##0"
        S.cell(row=rr, column=4, value=f'=IFERROR(AVERAGEIFS({rng("YT retention d7")},{key_range},{a},{rng("YT retention d7")},">0"),"")').number_format = "0.0%"
        S.cell(row=rr, column=5, value=f'=IFERROR(AVERAGEIFS({rng("Engagement d7")},{key_range},{a},{rng("Engagement d7")},">0"),"")').number_format = "0.0%"
        S.cell(row=rr, column=6, value=f'=IFERROR(AVERAGEIFS({rng("Total views d30")},{key_range},{a},{rng("Total views d30")},">0"),"")').number_format = "#,##0"
        for j in range(2, 7):
            S.cell(row=rr, column=j).font = Font(name="Arial", size=11)
    return top + len(keys) + 2

nxt = block(4, "Theme", f"{SH}!$B${R0}:$B${R1}", themes)
nxt = block(nxt, "Hook style", rng("Hook style"), ["eyes", "thoughts"])
nxt = block(nxt, "Weekday", rng("Weekday"), ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"])
block(nxt, "Template fit", rng("Template fit"), ["Strong", "Weak"])
for c, w in zip("ABCDEF", (38, 11, 13, 15, 17, 13)):
    S.column_dimensions[c].width = w
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(dst)
print("order:", order[:20], "...")
